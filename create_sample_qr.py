"""
QR / 바코드 생성기 테스트용 샘플 엑셀 생성
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
import random
import os

# ── 샘플 데이터 ──────────────────────────────────
categories = ["음료", "스낵", "유제품", "냉동식품", "생활용품"]
brands = {
    "음료":   ["코카콜라", "펩시", "스프라이트", "환타", "파워에이드"],
    "스낵":   ["포카칩", "새우깡", "꼬깔콘", "오레오", "프링글스"],
    "유제품": ["서울우유", "남양유업", "매일유업", "빙그레", "상하목장"],
    "냉동식품":["비비고 만두", "CJ 고메", "동원 참치", "오뚜기 카레", "풀무원 순두부"],
    "생활용품":["다우니", "피죤", "리엔", "케라시스", "헤드앤숄더"],
}

products = []
for cat, brand_list in brands.items():
    for brand in brand_list:
        # 바코드는 Code128 호환 — ASCII 영숫자만 사용
        cat_code = {"음료":"BV","스낵":"SN","유제품":"DY","냉동식품":"FZ","생활용품":"HH"}[cat]
        code128 = f"{cat_code}{str(random.randint(10000, 99999))}"
        # EAN-13 테스트용 (12자리 숫자, 체크섬은 JsBarcode가 자동 계산)
        ean13 = "880" + str(random.randint(100000000, 999999999))  # 12자리
        price = random.choice([980, 1200, 1500, 1800, 2500, 3200, 4500, 5900, 7800])
        stock = random.randint(5, 200)
        location = f"{random.choice(['A','B','C','D'])}-{random.randint(1,5):02d}-{random.randint(1,12):02d}"
        url = f"https://shop.example.com/product/{code128.lower()}"

        products.append({
            "상품코드":    code128,
            "EAN13코드":   ean13,
            "상품명":      brand,
            "카테고리":    cat,
            "단가":        price,
            "재고":        stock,
            "창고위치":    location,
            "상품URL":     url,
        })

random.shuffle(products)

# ── 엑셀 생성 ──────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "상품목록"

headers = ["상품코드", "EAN13코드", "상품명", "카테고리", "단가", "재고", "창고위치", "상품URL"]

# 헤더 스타일
hdr_fill  = PatternFill("solid", fgColor="1A3A2A")
hdr_font  = Font(name="맑은 고딕", bold=True, color="00C878", size=11)
hdr_align = Alignment(horizontal="center", vertical="center")
thin      = Side(style="thin", color="2B5C3E")
hdr_border= Border(bottom=Side(style="medium", color="00C878"))

ws.row_dimensions[1].height = 30

for ci, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=ci, value=h)
    cell.font   = hdr_font
    cell.fill   = hdr_fill
    cell.alignment = hdr_align
    cell.border = hdr_border

# 데이터 스타일
even_fill = PatternFill("solid", fgColor="F4F9F6")
odd_fill  = PatternFill("solid", fgColor="FFFFFF")
data_font = Font(name="맑은 고딕", size=10)
center    = Alignment(horizontal="center", vertical="center")
left      = Alignment(horizontal="left",   vertical="center")

for ri, row in enumerate(products, 2):
    fill = even_fill if ri % 2 == 0 else odd_fill
    ws.row_dimensions[ri].height = 22
    for ci, key in enumerate(headers, 1):
        val  = row[key]
        cell = ws.cell(row=ri, column=ci, value=val)
        cell.font   = data_font
        cell.fill   = fill
        cell.border = Border(
            bottom=Side(style="thin", color="DDDDDD"),
            right =Side(style="thin", color="EEEEEE"),
        )
        if key in ("단가", "재고"):
            cell.alignment = center
            if key == "단가":
                cell.number_format = '#,##0"원"'
        elif key in ("상품명", "상품URL", "창고위치"):
            cell.alignment = left
        else:
            cell.alignment = center

# 컬럼 너비
col_widths = [14, 16, 18, 12, 10, 8, 14, 42]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# 틀 고정
ws.freeze_panes = "A2"

# 자동 필터
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── 두 번째 시트: 직원 ID 카드용 ──────────────────
ws2 = wb.create_sheet("직원목록")

dept_list = ["개발팀", "디자인팀", "마케팅팀", "영업팀", "인사팀", "재무팀"]
names = [
    "김지윤","이준혁","박소연","최민준","정예린","강태양","윤서현","임도현",
    "한채원","오민서","신유진","류지호","배수아","전민재","황나연","남주원",
    "문세진","심은지","고현우","노아린","장하늘","엄지수","곽태민","원소희",
]
random.shuffle(names)

hdr2  = ["사원번호", "이름", "부서", "직급", "입사일", "이메일", "사원QR값"]
positions = ["사원", "대리", "과장", "차장", "부장"]

ws2.row_dimensions[1].height = 30
for ci, h in enumerate(hdr2, 1):
    cell = ws2.cell(row=1, column=ci, value=h)
    cell.font   = Font(name="맑은 고딕", bold=True, color="1A3A6A", size=11)
    cell.fill   = PatternFill("solid", fgColor="E8F0FE")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(bottom=Side(style="medium", color="4A6FA5"))

import datetime
for ri, name in enumerate(names, 2):
    dept  = random.choice(dept_list)
    pos   = random.choice(positions)
    emp_id= f"EMP{ri-1:04d}"
    year  = random.randint(2018, 2024)
    month = random.randint(1, 12)
    day   = random.randint(1, 28)
    hire  = datetime.date(year, month, day)
    email = f"{name.replace('','').lower()}{ri-1}@company.com"
    qr_val= f"EMP:{emp_id}|NAME:{name}|DEPT:{dept}"

    fill2 = PatternFill("solid", fgColor="F8FAFF") if ri % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
    for ci, val in enumerate([emp_id, name, dept, pos, hire, email, qr_val], 1):
        cell = ws2.cell(row=ri, column=ci, value=val)
        cell.font      = Font(name="맑은 고딕", size=10)
        cell.fill      = fill2
        cell.alignment = Alignment(horizontal="center" if ci != 6 else "left", vertical="center")
        cell.border    = Border(bottom=Side(style="thin", color="DDDDDD"))
        if ci == 5 and isinstance(val, datetime.date):
            cell.number_format = "YYYY-MM-DD"
    ws2.row_dimensions[ri].height = 22

col_widths2 = [12, 10, 12, 8, 13, 26, 36]
for i, w in enumerate(col_widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = f"A1:{get_column_letter(len(hdr2))}1"

# ── 저장 ──
out = os.path.join(os.path.dirname(__file__), "sample_qr_test.xlsx")
wb.save(out)
print(f"저장 완료: {out}")
print(f"  시트1 '상품목록': {len(products)}행")
print(f"  시트2 '직원목록': {len(names)}행")
